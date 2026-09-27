#!/usr/bin/env python3
"""Build the Close as a Quiet Prayer singalong deck.

Uses images/NN.png where present. Any slide whose image is missing is drawn
text-only per references/text-slide-design-spec.md, so the deck is usable
tonight and upgrades itself as images are committed.

    python3 build_deck.py            # illustrated where possible
    python3 build_deck.py --text     # force text-only for every slide
    python3 build_deck.py --png      # also export 1920x1080 PNGs
"""
import argparse, re, subprocess, sys, tempfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
TITLE = "Close as a Quiet Prayer"
CREDIT = "Words and music: Sally DeFord  ·  Hymns—For Home and Church, 1030"

# Slide position -> image number (title first). Refrain image 03 repeats.
IMAGE_FOR = ["11", "01", "02", "03", "04", "05", "06", "03", "07", "03", "08", "09", "10", "03"]

# ---- DESIGN SPEC -----------------------------------------------------------
W_IN, H_IN = 13.333, 7.5
RENDER = (2000, 1125)
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
# Illustrated (matches Holding Hands): bottom scrim, bottom-anchored lyrics
SCRIM_FRAC, SCRIM_ALPHA, SCRIM_EXP = 0.40, 240, 1.15
ILL_BOX = (0.55, 5.20, 12.233, 1.95)       # x, y, w, h in inches
ILL_MAX_PT, ILL_MIN_PT = 40, 32
ILL_USABLE_IN = 11.9
# Text-only (references/text-slide-design-spec.md)
NAVY_TOP, NAVY_BOT, LIFT = (0x1C, 0x2E, 0x4E), (0x0A, 0x11, 0x20), (0x2E, 0x46, 0x70)
TITLE_TOP = (0x24, 0x38, 0x5C)
CREAM, GOLD, CREDIT_C = RGBColor(0xFD, 0xFA, 0xEE), RGBColor(0xD9, 0xB8, 0x6A), RGBColor(0xA8, 0xB4, 0xC8)
TXT_BOX = (0.667, 1.30, 12.0, 5.2)
TXT_USABLE_IN, TXT_RANGE, TXT_FLOOR = 11.6, (28, 60), 44
LINE_SPACING = 1.35
# ---------------------------------------------------------------------------

CLAUSE_BEFORE = {"and", "or", "when", "but"}
BREAK_BEFORE = {"of", "in", "to", "with", "that", "who", "around", "for", "by", "anytime,"}


def width_in(text, pt):
    return ImageFont.truetype(FONT_BOLD, pt * 10).getlength(text) / 10 / 72


def smart_break(line, pt, usable):
    """Split one overflowing line at a phrase boundary, as late as fits."""
    words = line.split()
    cands = []
    for i in range(1, len(words)):
        a, b = " ".join(words[:i]), " ".join(words[i:])
        if width_in(a, pt) > usable or len(words) - i < 2:
            continue
        tier = 0 if (words[i - 1][-1] in ",;:" or words[i].lower() in CLAUSE_BEFORE) else 1 if words[i].lower() in BREAK_BEFORE else 2
        cands.append((tier, -i, [a, b]))
    return min(cands)[2] if cands else [line]


def parse_lyrics():
    text = (HERE / "lyrics.md").read_text(encoding="utf-8")
    body = text[text.index("## Verse 1"):]
    slides = []
    for block in re.split(r"\n\s*\n", body):
        lines = [l.strip() for l in block.splitlines() if l.strip() and not l.startswith("#")]
        if lines:
            slides.append(lines)
    if len(slides) != len(IMAGE_FOR) - 1:
        sys.exit(f"expected {len(IMAGE_FOR) - 1} lyric slides, found {len(slides)}")
    return slides


def fit_lines(slides, usable, lo, hi, floor=None):
    """One size for the deck; break only lines that overflow at the floor."""
    def best(ss):
        for pt in range(hi, lo - 1, -1):
            if all(width_in(l, pt) <= usable for s in ss for l in s):
                return pt
        return lo
    pt = best(slides)
    if pt < (floor or lo):
        f = floor or lo
        slides = [[p for l in s for p in (smart_break(l, f, usable) if width_in(l, f) > usable else [l])] for s in slides]
        pt = best(slides)
    return pt, slides


def cover(path):
    im = Image.open(path).convert("RGB")
    r = im.width / im.height
    if abs(r - 16 / 9) > 0.01:
        print(f"  warning: {path.name} is {im.width}x{im.height} ({r:.3f}), not 16:9; center-cropping")
    tw, th = (im.width, round(im.width * 9 / 16)) if r < 16 / 9 else (round(im.height * 16 / 9), im.height)
    l, t = (im.width - tw) // 2, (im.height - th) // 2
    return im.crop((l, t, l + tw, t + th)).resize(RENDER, Image.LANCZOS)


def scrimmed(path):
    im = cover(path).convert("RGBA")
    w, h = RENDER
    band = int(h * SCRIM_FRAC)
    grad = Image.new("L", (1, band))
    for y in range(band):
        grad.putpixel((0, y), int(SCRIM_ALPHA * (y / (band - 1)) ** SCRIM_EXP))
    over = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    over.paste(Image.new("RGBA", (w, band), (0, 0, 0, 255)), (0, h - band), grad.resize((w, band)))
    return Image.alpha_composite(im, over).convert("RGB")


def navy(top):
    w, h = RENDER
    im = Image.new("RGB", RENDER)
    d = ImageDraw.Draw(im)
    for y in range(h):
        t = y / (h - 1)
        d.line([(0, y), (w, y)], fill=tuple(round(a + (b - a) * t) for a, b in zip(top, NAVY_BOT)))
    glow = Image.new("L", RENDER, 0)
    ImageDraw.Draw(glow).ellipse([w * 0.18, h * 0.44 - h * 0.30, w * 0.82, h * 0.44 + h * 0.30], fill=150)
    glow = glow.filter(ImageFilter.GaussianBlur(160))
    return Image.composite(Image.new("RGB", RENDER, LIFT), im, glow)


def add_bg(slide, img, tmp, name):
    p = Path(tmp) / f"{name}.png"
    img.save(p)
    slide.shapes.add_picture(str(p), 0, 0, Inches(W_IN), Inches(H_IN))


def add_text(slide, box, lines, pt, color, anchor, spacing=1.0, size_map=None):
    x, y, w, h = box
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment, para.line_spacing = PP_ALIGN.CENTER, spacing
        run = para.add_run()
        run.text = line
        f = run.font
        f.name, f.bold, f.size, f.color.rgb = "Arial", (size_map or {}).get(i, (True,))[0], Pt((size_map or {}).get(i, (True, pt))[1]), color
        if size_map and i in size_map and len(size_map[i]) > 2:
            f.color.rgb = size_map[i][2]
    return tf


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--text", action="store_true")
    ap.add_argument("--png", action="store_true")
    ap.add_argument("-o", default=str(HERE / "slideshow.pptx"))
    a = ap.parse_args()

    slides = parse_lyrics()
    ill_pt, ill_slides = fit_lines(slides, ILL_USABLE_IN, ILL_MIN_PT, ILL_MAX_PT, floor=ILL_MAX_PT)
    txt_pt, txt_slides = fit_lines(slides, TXT_USABLE_IN, *TXT_RANGE, floor=TXT_FLOOR)
    have = {n: (HERE / "images" / f"{n}.png") for n in set(IMAGE_FOR)}
    have = {n: p for n, p in have.items() if p.exists() and not a.text}
    print(f"illustrated lyric size {ill_pt} pt; text-only size {txt_pt} pt; images present: {sorted(have) or 'none'}")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    blank = prs.slide_layouts[6]
    names = []
    with tempfile.TemporaryDirectory() as tmp:
        # Title
        s = prs.slides.add_slide(blank)
        if "11" in have:
            add_bg(s, scrimmed(have["11"]), tmp, "t")
            add_text(s, ILL_BOX, [TITLE, CREDIT], 48, RGBColor(255, 255, 255), MSO_ANCHOR.BOTTOM,
                     1.1, {0: (True, 48), 1: (False, 16, RGBColor(0xE8, 0xE8, 0xE8))})
        else:
            add_bg(s, navy(TITLE_TOP), tmp, "t")
            add_text(s, (0.667, 2.2, 12.0, 1.4), [TITLE], 52, CREAM, MSO_ANCHOR.BOTTOM)
            rule = s.shapes.add_shape(1, Inches((W_IN - 2.2) / 2), Inches(3.8), Inches(2.2), Pt(2.5))
            rule.fill.solid(); rule.fill.fore_color.rgb = GOLD; rule.line.fill.background()
            add_text(s, (0.667, 4.1, 12.0, 0.6), [CREDIT], 16, CREDIT_C, MSO_ANCHOR.TOP, 1.0, {0: (False, 16)})
        names.append("00 - Title")
        # Lyrics
        for k, n in enumerate(IMAGE_FOR[1:]):
            s = prs.slides.add_slide(blank)
            if n in have:
                add_bg(s, scrimmed(have[n]), tmp, f"s{k}")
                add_text(s, ILL_BOX, ill_slides[k], ill_pt, RGBColor(255, 255, 255), MSO_ANCHOR.BOTTOM, 1.1)
            else:
                add_bg(s, navy(NAVY_TOP), tmp, f"s{k}")
                add_text(s, TXT_BOX, txt_slides[k], txt_pt, CREAM, MSO_ANCHOR.MIDDLE, LINE_SPACING)
            clean = re.sub(r'[\\/:*?"<>|]', "", slides[k][0]).rstrip(",.;:")
            names.append(f"{k + 1:02d} - {clean}")
        prs.save(a.o)
    print(f"saved {a.o} ({len(names)} slides)")

    if a.png:
        out = Path(a.o).with_suffix("")
        out = out.parent / (out.name + " - PNG")
        out.mkdir(exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), a.o], check=True, capture_output=True)
        pdf = out / (Path(a.o).stem + ".pdf")
        subprocess.run(["pdftoppm", "-png", "-r", "144", str(pdf), str(out / "p")], check=True)
        pngs = sorted(out.glob("p-*.png"))
        for p, nm in zip(pngs, names):
            p.rename(out / f"{nm}.png")
        pdf.unlink()
        print(f"exported {len(pngs)} PNGs to {out}")


if __name__ == "__main__":
    main()
