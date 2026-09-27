#!/usr/bin/env python3
"""Grid view: the finished slides at 1/4 size, up to four per page, on black.

One page per song section, so sections never share a page. The title stays
full size. Incomplete rows are centered. Tiles are the main deck's rendered
1920x1080 slides, so they match it exactly.

    python3 build_deck.py --png && python3 build_grid.py
"""
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

HERE = Path(__file__).resolve().parent
PNG_DIR = HERE / "slideshow - PNG"
OUT = HERE / "slideshow - grid.pptx"
W, H = 13.333, 7.5
TW, TH = W / 2, H / 2

# Main-deck slide numbers (00 = title) grouped by section.
PAGES = [
    ("Title", [0]),
    ("Verse 1", [1, 2, 3]),
    ("Verse 2", [4, 5, 6, 7]),
    ("Bridge", [8, 9]),
    ("Verse 3", [10, 11, 12, 13]),
]


def positions(n):
    """Top row then bottom row; a short row is centered; two tiles sit mid-slide."""
    if n == 1:
        return [(0, 0, W, H)]
    if n == 2:
        return [(0, (H - TH) / 2, TW, TH), (TW, (H - TH) / 2, TW, TH)]
    top = [(0, 0, TW, TH), (TW, 0, TW, TH)]
    bottom = [(0, TH, TW, TH), (TW, TH, TW, TH)] if n == 4 else [((W - TW) / 2, TH, TW, TH)]
    return top + bottom


def main():
    pngs = sorted(PNG_DIR.glob("*.png"))
    if len(pngs) != 14:
        raise SystemExit(f"expected 14 rendered slides in {PNG_DIR}, found {len(pngs)}; run build_deck.py --png first")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    with tempfile.TemporaryDirectory() as tmp:
        for name, idx in PAGES:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
            for i, (x, y, w, h) in zip(idx, positions(len(idx))):
                jpg = Path(tmp) / f"{i:02d}.jpg"
                Image.open(pngs[i]).convert("RGB").save(jpg, quality=90, subsampling=0)
                s.shapes.add_picture(str(jpg), Inches(x), Inches(y), Inches(w), Inches(h))
            s.notes_slide.notes_text_frame.text = name
        prs.save(OUT)
    print(f"saved {OUT} ({len(PAGES)} slides)")


if __name__ == "__main__":
    main()
